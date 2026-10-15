"""Bounded, validated Excel imports and structured Excel exports."""
from __future__ import annotations

import io
import unicodedata
from datetime import datetime, time
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.auditlog.models import AuditEvent
from apps.auditlog.services import record_event
from apps.accounts.models import User
from apps.operations.models import CollectionProduct, FiscalParameter, MilkReceipt, Producer, ProducerProduct, Route, Vehicle
from .documents import create_receipt


SHEETS = {"routes": "Vehículos", "producers": "Productores", "vehicles": "Vehículos", "receipts": "Recolecciones"}
HEADERS = {
    "routes": ["Ruta", "Descripción", "Estado"],
    "producers": ["Tipo de documento", "Número de documento / NIT", "Nombre / Razón social", "Nombre de la finca", "Teléfono de contacto", "Municipio", "VEREDA", "Observaciones", "Estado", "Clasificación", "Productos (códigos separados por |)", "Retenciones (nombres separados por |)", "ID proveedor"],
    "vehicles": ["Nombre / Placa del vehículo", "Ruta", "Nombre del conductor", "Documento del conductor", "Teléfono del conductor", "Capacidad del tanque (litros)", "Observaciones", "Usuario conductor", "Nombre del vehículo", "Estado"],
    "receipts": ["RUTA", "ID / No. de recolección", "Productor", "Vehículo / Placa", "Fecha", "Hora", "Litros recogidos", "Estado de recepción", "Soporte", "Medio de pago", "Comprobante / No. soporte", "Observaciones", "Conductor del recorrido", "Producto", "Precio histórico COP/L"],
}


def norm(value):
    return " ".join(unicodedata.normalize("NFKD", str(value or "").strip()).encode("ascii", "ignore").decode().lower().split())


def cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def required(value, label):
    value = cell(value)
    if not value:
        raise ValueError(f"Falta {label}.")
    return value


def decimal(value):
    try:
        result = Decimal(cell(value).replace(",", "."))
        if result <= 0:
            raise ValueError
        return result.quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        raise ValueError("La capacidad o los litros deben ser un número positivo.") from None


def with_active(row, value):
    if cell(value):
        state = norm(value)
        if state not in {"activo", "activa", "inactivo", "inactiva", "si", "no", "true", "false"}:
            raise ValueError("Estado inválido; use Activo o Inactivo.")
        row["active"] = state in {"activo", "activa", "si", "true"}
    return row


def parse_row(kind, data):
    get = lambda index: data.get(norm(HEADERS[kind][index]), "")
    if kind == "routes":
        return with_active({"name": required(get(0), "la ruta").upper(), "description": cell(get(1))}, get(2))
    if kind == "producers":
        category = {"cuantias menores": "SMALL", "small": "SMALL", "documento soporte": "SUPPORT", "support": "SUPPORT", "facturacion electronica": "ELECTRONIC", "electronic": "ELECTRONIC", "": "SUPPORT"}.get(norm(get(9)))
        if not category:
            raise ValueError("Clasificación desconocida.")
        dtype = norm(get(0))
        dtype = {"cedula": "CC", "cedula de ciudadania": "CC", "cc": "CC", "nit": "NIT", "ce": "CE", "cedula de extranjeria": "CE", "otro": "OTHER"}.get(dtype)
        if not dtype and (category != "SMALL" or cell(get(0))):
            raise ValueError("Tipo de documento desconocido.")
        row = with_active({"fiscal_category": category, "document_type": dtype or "", "document_number": "".join((cell(get(1)) if category == "SMALL" else required(get(1), "el documento")).split()).upper(), "name": required(get(2), "el nombre"), "farm_name": cell(get(3)), "phone": cell(get(4)), "municipality": cell(get(5)), "village": cell(get(6)), "notes": cell(get(7))}, get(8))
        row["_provider_id"] = cell(get(12))
        row["_products"] = cell(get(10))
        row["_fiscal"] = cell(get(11)) if norm(HEADERS[kind][11]) in data else None
        row["_prices"] = {p.code: data.get(norm(f"Precio {p.code} COP/L"), "") for p in CollectionProduct.objects.all() if norm(f"Precio {p.code} COP/L") in data}
        return row
    if kind == "vehicles":
        return with_active({"plate": "".join(required(get(0), "la placa").split()).upper(), "route_name": cell(get(1)).upper(), "driver_name": cell(get(2)), "driver_document": cell(get(3)), "driver_phone": cell(get(4)), "capacity_liters": decimal(get(5)) if cell(get(5)) else None, "notes": cell(get(6)), "driver_username": cell(get(7)), "name": cell(get(8))}, get(9))
    reference = required(get(1), "el ID de recolección")
    producer = required(get(2), "el productor")
    plate = required(get(3), "la placa").upper()
    date_value, time_value = get(4), get(5)
    if isinstance(date_value, datetime):
        day = date_value.date()
    else:
        try:
            day = datetime.strptime(cell(date_value), "%Y-%m-%d").date()
        except ValueError:
            raise ValueError("Fecha inválida; use AAAA-MM-DD.") from None
    if isinstance(time_value, (datetime, time)):
        clock = time_value.time() if isinstance(time_value, datetime) else time_value
    else:
        try:
            clock = datetime.strptime(cell(time_value), "%H:%M").time()
        except ValueError:
            raise ValueError("Hora inválida; use HH:MM, no solo AM/PM.") from None
    if norm(get(7)) not in {"confirmado", "confirmed", "recibido"}:
        raise ValueError("Solo se importan recibidos confirmados; las anulaciones se gestionan en la plataforma.")
    return {"source_reference": reference, "producer_key": producer, "plate": plate, "route_name": cell(get(0)).upper(), "collected_at": timezone.make_aware(datetime.combine(day, clock), timezone.get_current_timezone()), "liters": decimal(get(6)), "support_type": cell(get(8)), "payment_method": cell(get(9)), "support_number": cell(get(10)), "observations": cell(get(11))}


def read_rows(upload, kind):
    try:
        book = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        raise ValueError("No se pudo leer el archivo Excel.") from None
    try:
        sheet_name = "Rutas" if kind == "routes" and "Rutas" in book.sheetnames else SHEETS[kind]
        if sheet_name not in book.sheetnames:
            raise ValueError(f"Falta la hoja «{sheet_name}».")
        sheet = book[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("La hoja está vacía.")
        index = {norm(value): position for position, value in enumerate(headers) if value}
        needed = HEADERS[kind][:1] if kind == "routes" else HEADERS[kind][:8] if kind == "producers" else HEADERS[kind][:7] if kind == "vehicles" else HEADERS[kind][:12]
        missing = [h for h in needed if norm(h) not in index]
        if missing:
            raise ValueError("Faltan columnas: " + ", ".join(missing))
        data_rows = []
        errors = []
        seen = set()
        for number, values in enumerate(rows, 2):
            if number > 5002:
                raise ValueError("Máximo 5000 filas por importación.")
            if not any(cell(v) for v in values):
                continue
            raw = {key: values[pos] if pos < len(values) else "" for key, pos in index.items()}
            if kind == "routes":
                raw[norm("Descripción")] = raw.get(norm("Descripción"), "")
                if not cell(raw.get(norm("Ruta"))):
                    continue
            try:
                parsed = parse_row(kind, raw)
                if kind == "receipts":
                    code = cell(raw.get(norm("Producto"))) or "leche"
                    product = CollectionProduct.objects.filter(code__iexact=code, active=True).first()
                    if not product:
                        raise ValueError("Producto inexistente o inactivo. Use el código del catálogo.")
                    parsed["product"] = product
                limits = {"name": 180, "description": 300, "farm_name": 180, "plate": 12, "document_number": 30, "phone": 30, "municipality": 120, "village": 120, "driver_name": 180, "driver_document": 30, "driver_phone": 30, "driver_username": 150, "source_reference": 80, "support_type": 80, "payment_method": 80, "support_number": 80, "notes": 1000, "observations": 1000}
                if kind == "routes":
                    limits["name"] = 120
                for field, maximum in limits.items():
                    if len(cell(parsed.get(field))) > maximum:
                        raise ValueError(f"El campo {field} supera {maximum} caracteres.")
                if kind == "producers":
                    key = parsed["_provider_id"] or ((parsed["document_type"], parsed["document_number"]) if parsed["document_number"] else ("SMALL", norm(parsed["name"])))
                elif kind == "receipts":
                    key = parsed["source_reference"]
                elif kind == "vehicles":
                    key = parsed["plate"]
                else:
                    key = parsed["name"]
                if key in seen:
                    if kind == "routes":
                        continue
                    raise ValueError("Identificador repetido en el archivo.")
                seen.add(key)
                data_rows.append((number, parsed))
            except ValueError as exc:
                errors.append(f"Fila {number}: {exc}")
        if not data_rows and not errors:
            errors.append("No hay filas para importar.")
        return data_rows, errors
    finally:
        book.close()


@transaction.atomic
def import_workbook(upload, kind, user, request=None):
    rows, errors = read_rows(upload, kind)
    prepared = []
    if kind == "receipts":
        for number, row in rows:
            candidates = Producer.objects.filter(document_number__iexact=row["producer_key"])
            if not candidates.exists():
                candidates = Producer.objects.filter(name__iexact=row["producer_key"])
            if candidates.count() != 1:
                errors.append(f"Fila {number}: productor ambiguo o inexistente; use su documento o nombre exacto.")
            if not Vehicle.objects.filter(plate__iexact=row["plate"], active=True).exists():
                errors.append(f"Fila {number}: vehículo inexistente o inactivo ({row['plate']}).")
            if row["route_name"] and not Route.objects.filter(name__iexact=row["route_name"], active=True).exists():
                errors.append(f"Fila {number}: ruta inexistente ({row['route_name']}).")
            if MilkReceipt.objects.filter(source_reference=row["source_reference"]).exists():
                errors.append(f"Fila {number}: ID de recolección ya registrado.")
            prepared.append((row, candidates.first()))
    if kind == "vehicles":
        for number, row in rows:
            username = row["driver_username"]
            if username and not User.objects.filter(username=username, is_active=True, role__in=[User.Role.DRIVER, User.Role.COLLECTOR]).exists():
                errors.append(f"Fila {number}: usuario conductor inexistente o sin rol de conductor/recolector ({username}).")
    if errors:
        return 0, 0, errors[:50]
    created = updated = 0
    for index, (number, row) in enumerate(rows):
        if kind == "producers":
            is_created = import_producer_row(row, user)
            created += is_created
            updated += not is_created
            continue
        if kind == "receipts":
            producer = prepared[index][1]
            vehicle = Vehicle.objects.get(plate__iexact=row.pop("plate"))
            route_name = row.pop("route_name")
            row.pop("producer_key")
            create_receipt(user=user, request=request, producer=producer, vehicle=vehicle, route=Route.objects.get(name__iexact=route_name) if route_name else vehicle.route, source=MilkReceipt.Source.IMPORT, **row)
            created += 1
            continue
        if kind == "routes":
            model, lookup = Route, {"name": row.pop("name")}
        elif kind == "producers":
            model, lookup = Producer, {"document_type": row.pop("document_type"), "document_number": row.pop("document_number")}
        else:
            model, lookup = Vehicle, {"plate": row.pop("plate")}
            name = row.pop("route_name")
            username = row.pop("driver_username")
            row["route"] = Route.objects.get_or_create(name=name)[0] if name else None
            if username:
                row["driver"] = User.objects.get(username=username)
        obj, is_created = model.objects.update_or_create(**lookup, defaults=row)
        record_event(actor=user, action=AuditEvent.Action.CREATE if is_created else AuditEvent.Action.UPDATE, instance=obj, changes={"imported_from": SHEETS[kind], "row": number}, request=request)
        created += is_created
        updated += not is_created
    return created, updated, []


def safe_excel(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def workbook_response(book, filename):
    buffer = io.BytesIO()
    book.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def make_sheet(book, title, headers, records):
    sheet = book.create_sheet(title)
    sheet.append(headers)
    for record in records:
        sheet.append([safe_excel(value) for value in record])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for c in sheet[1]:
        c.fill = PatternFill("solid", fgColor="174B4C")
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        sheet.column_dimensions[letter].width = min(45, max(15, max(len(str(c.value or "")) for c in column[:200]) + 2))
    return sheet


def export_workbook(kind):
    book = Workbook()
    book.remove(book.active)
    if kind == "routes":
        records = ([o.name, o.description, "Activo" if o.active else "Inactivo"] for o in Route.objects.all())
    elif kind == "producers":
        products = list(CollectionProduct.objects.all())
        records = ([o.get_document_type_display(), o.document_number, o.name, o.farm_name, o.phone, o.municipality, o.village, o.notes, "Activo" if o.active else "Inactivo", o.get_fiscal_category_display(), "|".join(o.products.values_list("code", flat=True)), "|".join(o.deductions.values_list("name", flat=True)), str(o.pk), *[(o.product_prices.filter(product=p).first().unit_price if o.product_prices.filter(product=p).exists() else None) for p in products]] for o in Producer.objects.all())
    elif kind == "vehicles":
        records = ([o.plate, o.route.name if o.route else "", o.driver_name, o.driver_document, o.driver_phone, o.capacity_liters or "", o.notes, o.driver.username if o.driver else "", o.name, "Activo" if o.active else "Inactivo"] for o in Vehicle.objects.select_related("route", "driver"))
    else:
        records = ([o.route.name if o.route else "", o.source_reference or o.internal_number, o.producer.document_number, o.vehicle.plate, timezone.localtime(o.collected_at).date().isoformat(), timezone.localtime(o.collected_at).strftime("%H:%M"), o.liters, o.get_status_display(), o.support_type, o.payment_method, o.support_number, o.observations, (o.trip.driver.get_full_name() or o.trip.driver.username) if o.trip else ""] for o in MilkReceipt.objects.select_related("route", "producer", "vehicle", "trip__driver"))
    headers = HEADERS[kind] + ([f"Precio {p.code} COP/L" for p in products] if kind == "producers" else [])
    if kind == "receipts":
        records = ([o.route.name if o.route else "", o.source_reference or o.internal_number, o.producer.document_number or o.producer.name, o.vehicle.plate, timezone.localtime(o.collected_at).date().isoformat(), timezone.localtime(o.collected_at).strftime("%H:%M"), o.liters, o.get_status_display(), o.support_type, o.payment_method, o.support_number, o.observations, (o.trip.driver.get_full_name() or o.trip.driver.username) if o.trip else "", o.product.code, o.unit_price] for o in MilkReceipt.objects.select_related("route", "producer", "product", "vehicle", "trip__driver"))
    make_sheet(book, "Rutas" if kind == "routes" else SHEETS[kind], headers, records)
    return workbook_response(book, f"ubalac_{kind}")


def import_producer_row(row, user):
    import uuid
    from .pricing import snapshot_terms
    provider_id, codes, fiscal, prices = (row.pop(k) for k in ("_provider_id", "_products", "_fiscal", "_prices"))
    if provider_id:
        try:
            pk = uuid.UUID(provider_id)
        except ValueError:
            raise ValidationError("ID de proveedor inválido.")
        obj = Producer.objects.filter(pk=pk).first()
        if not obj:
            raise ValidationError("ID de proveedor inexistente; déjelo vacío para crear.")
    elif row["document_number"]:
        obj = Producer.objects.filter(document_type=row["document_type"], document_number=row["document_number"]).first()
    else:
        matches = Producer.objects.filter(name__iexact=row["name"], fiscal_category="SMALL")
        if matches.count() > 1:
            raise ValidationError("Proveedor ambiguo; use el ID proveedor de la exportación.")
        obj = matches.first()
    is_created = obj is None
    obj = obj or Producer()
    for k, value in row.items():
        setattr(obj, k, value)
    obj.full_clean()
    obj.save()
    if fiscal is not None:
        if not (user.is_superuser or user.role == "ADMIN"):
            names = {n.strip() for n in fiscal.split("|") if n.strip()}
            if names != set(obj.deductions.values_list("name", flat=True)):
                raise ValidationError("Solo el administrador puede asignar retenciones desde Excel.")
        else:
            names = [n.strip() for n in fiscal.split("|") if n.strip()]
            parameters = FiscalParameter.objects.filter(name__in=names)
            if parameters.count() != len(set(names)):
                raise ValidationError("Retención desconocida. Créela primero en el módulo de parámetros.")
            if sum((p.rate for p in parameters if p.kind == "WITHHOLDING" and p.active), Decimal("0")) > 100:
                raise ValidationError("La suma de retenciones supera 100 %.")
            obj.deductions.set(parameters)
    if codes:
        codes = [c.strip() for c in codes.split("|") if c.strip()]
        selected = list(CollectionProduct.objects.filter(code__in=codes, active=True))
        if len(selected) != len(set(codes)):
            raise ValidationError("Producto inexistente o inactivo.")
        obj.product_prices.exclude(product__in=selected).delete()
    elif obj.product_prices.exists():
        selected = list(obj.products.all())
    else:
        selected = list(CollectionProduct.objects.filter(code="leche"))
    for product in selected:
        price, _ = ProducerProduct.objects.get_or_create(producer=obj, product=product)
        if product.code in prices and cell(prices[product.code]) != "":
            try:
                price.unit_price = Decimal(cell(prices[product.code]).replace(",", "."))
            except InvalidOperation:
                raise ValidationError("Precio inválido.") from None
        price.full_clean()
        price.save()
    snapshot_terms(obj)
    record_event(actor=user, action=AuditEvent.Action.UPDATE, instance=obj, changes={"imported_from": "Productores"})
    return is_created
