import io
from datetime import timedelta
from datetime import datetime, time

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.operations.models import MilkReceipt, Route, RouteTrip, default_product
from apps.operations.services import create_receipt
from apps.operations.services.trips import auto_close_expired_trips
from .factories import make_producer, make_user, make_vehicle, make_inspection


class TripAndReportTests(TestCase):
    def setUp(self):
        self.driver = make_user(username="driver-1", role=User.Role.DRIVER)
        self.other = make_user(username="driver-2", role=User.Role.DRIVER)
        self.viewer = make_user(username="viewer-report", role=User.Role.VIEWER)
        self.route = Route.objects.create(name="Norte")
        self.vehicle = make_vehicle()
        self.vehicle.driver = self.driver
        self.vehicle.route = self.route
        self.vehicle.save()
        self.producer = make_producer()
        self.inspection = make_inspection(self.driver, self.vehicle)

    def test_driver_opens_collects_and_closes_assigned_trip(self):
        self.client.force_login(self.driver)
        url = reverse("operations:trip_start")
        self.assertEqual(self.client.post(url, {"route": self.route.pk, "vehicle": self.vehicle.pk, "product": default_product(), "inspection": self.inspection.pk}).status_code, 302)
        trip = RouteTrip.objects.get()
        self.assertIsNone(trip.ended_at)
        response = self.client.post(url, {"route": self.route.pk, "vehicle": self.vehicle.pk, "product": default_product(), "inspection": self.inspection.pk})
        self.assertContains(response, "Ya tiene un recorrido abierto")
        receipt = create_receipt(user=self.driver, producer=self.producer, vehicle=self.vehicle, collected_at=timezone.now(), liters="55.25")
        self.assertEqual(receipt.trip_id, trip.pk)
        self.assertEqual(receipt.route_id, self.route.pk)
        self.client.force_login(self.other)
        self.assertEqual(self.client.post(reverse("operations:trip_close", kwargs={"pk": trip.pk})).status_code, 403)
        self.client.force_login(self.driver)
        self.assertEqual(self.client.post(reverse("operations:trip_close", kwargs={"pk": trip.pk})).status_code, 302)
        trip.refresh_from_db()
        self.assertIsNotNone(trip.ended_at)

    def test_driver_cannot_open_unassigned_vehicle_or_collect_without_trip(self):
        self.client.force_login(self.other)
        response = self.client.post(reverse("operations:trip_start"), {"route": self.route.pk, "vehicle": self.vehicle.pk, "product": default_product(), "inspection": self.inspection.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(RouteTrip.objects.count(), 0)
        response = self.client.post(reverse("operations:receipt_create"), {"producer": self.producer.pk, "vehicle": self.vehicle.pk, "collected_at": "2026-09-28T10:30", "liters": "5"})
        self.assertContains(response, "Escoja una opción válida")
        self.assertEqual(MilkReceipt.objects.count(), 0)

    def test_report_filters_and_excel_sheets(self):
        self.client.force_login(self.driver)
        self.client.post(reverse("operations:trip_start"), {"route": self.route.pk, "vehicle": self.vehicle.pk, "product": default_product(), "inspection": self.inspection.pk})
        trip = RouteTrip.objects.get()
        create_receipt(user=self.driver, producer=self.producer, vehicle=self.vehicle, collected_at=timezone.now(), liters="55.25")
        self.client.post(reverse("operations:trip_close", kwargs={"pk": trip.pk}))
        today = timezone.localdate().isoformat()
        params = {"date_from": today, "date_to": today, "producer": str(self.producer.pk), "route": str(self.route.pk), "driver": str(self.driver.pk)}
        response = self.client.get(reverse("dashboard:reports"), params)
        self.assertContains(response, "55,25")
        self.assertContains(response, "Desempeño de conductores")
        excel = self.client.get(reverse("dashboard:report_export"), params)
        book = load_workbook(io.BytesIO(excel.content), read_only=True, data_only=True)
        self.assertEqual(book.sheetnames, ["Liquidación por proveedor", "Retenciones e impuestos", "Detalle económico", "Resumen", "Productores", "Conductores", "Recibidos", "Recorridos", "Báscula"])
        self.assertEqual(book["Productores"]["C2"].value, 55.25)
        self.assertEqual(book["Conductores"]["B2"].value, 1)
        self.assertEqual(book["Recorridos"]["A2"].value, "NORTE")

    def test_closed_trip_without_receipts_appears_in_default_report_and_dashboard(self):
        self.client.force_login(self.driver)
        now = timezone.now()
        RouteTrip.objects.create(route=self.route, vehicle=self.vehicle, driver=self.driver, started_at=now - timedelta(minutes=5), ended_at=now)
        report = self.client.get(reverse("dashboard:reports"))
        self.assertContains(report, "Desempeño de conductores")
        self.assertContains(report, "driver-1")
        self.assertContains(report, "<strong>1</strong>", html=False)
        self.assertContains(report, 'class="report-sidebar"')
        dashboard = self.client.get(reverse("dashboard:index"))
        self.assertContains(dashboard, "Recorridos recientes")
        self.assertContains(dashboard, "NORTE")

    def test_overnight_trip_uses_closing_date(self):
        from apps.dashboard.reports import report_data

        now = timezone.now()
        RouteTrip.objects.create(route=self.route, vehicle=self.vehicle, driver=self.driver, started_at=now - timedelta(days=1, minutes=5), ended_at=now)
        data = report_data({"start": timezone.localdate(), "end": timezone.localdate()})
        self.assertEqual(data["completed_count"], 1)
        self.assertEqual(data["driver_rank"][0]["count"], 1)

    def test_daily_bars_scale_to_largest_day(self):
        from apps.dashboard.reports import report_data

        now = timezone.now()
        collector = make_user(username="collector-bars", role=User.Role.COLLECTOR)
        create_receipt(user=collector, producer=self.producer, vehicle=self.vehicle, collected_at=now - timedelta(days=1), liters="30")
        create_receipt(user=collector, producer=self.producer, vehicle=self.vehicle, collected_at=now, liters="90")
        data = report_data({"start": timezone.localdate() - timedelta(days=1), "end": timezone.localdate()})
        self.assertEqual([entry["bar_width"] for entry in data["daily"]], [33, 100])
        self.client.force_login(self.driver)
        page = self.client.get(reverse("dashboard:reports"), {"date_from": (timezone.localdate() - timedelta(days=6)).isoformat(), "date_to": timezone.localdate().isoformat()})
        self.assertContains(page, 'value="30.00"')
        self.assertContains(page, 'value="90.00"')

    def test_report_requires_both_dates_and_export_rejects_missing_date(self):
        self.client.force_login(self.driver)
        page = self.client.get(reverse("dashboard:reports"), {"date_from": timezone.localdate().isoformat()})
        self.assertContains(page, "Este campo es requerido")
        self.assertNotContains(page, "Descargar informe Excel")
        self.assertEqual(self.client.get(reverse("dashboard:report_export"), {"date_from": timezone.localdate().isoformat()}).status_code, 400)

    def test_api_driver_starts_and_closes_and_viewer_is_blocked(self):
        api = APIClient()
        api.force_authenticate(self.viewer)
        url = reverse("api-v1:route-trip-list")
        self.assertEqual(api.post(url, {"route": str(self.route.pk), "vehicle": str(self.vehicle.pk), "inspection": str(self.inspection.pk)}, format="json").status_code, 403)
        api.force_authenticate(self.driver)
        created = api.post(url, {"route": str(self.route.pk), "vehicle": str(self.vehicle.pk), "inspection": str(self.inspection.pk)}, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        closed = api.post(reverse("api-v1:route-trip-close", kwargs={"pk": created.data["id"]}))
        self.assertEqual(closed.status_code, 200, closed.data)
        self.assertIsNotNone(closed.data["ended_at"])

    def test_auto_close_at_2359_and_late_sync_of_offline_receipt(self):
        yesterday = timezone.localdate() - timedelta(days=1)
        started = timezone.make_aware(datetime.combine(yesterday, time(8, 0)))
        trip = RouteTrip.objects.create(route=self.route, vehicle=self.vehicle, driver=self.driver, started_at=started)
        self.assertEqual(auto_close_expired_trips(), 1)
        trip.refresh_from_db()
        self.assertEqual(timezone.localtime(trip.ended_at).time(), time(23, 59))
        captured = started + timedelta(minutes=18)
        receipt = create_receipt(user=self.driver, producer=self.producer, vehicle=self.vehicle, collected_at=captured, liters="18.50")
        self.assertEqual(receipt.trip_id, trip.pk)
        self.assertEqual(auto_close_expired_trips(), 0)
        api = APIClient()
        api.force_authenticate(self.driver)
        self.assertEqual(api.post(reverse("api-v1:route-trip-close", kwargs={"pk": trip.pk})).status_code, 200)
