import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.accounts.models import User
from apps.operations.models import MilkReceipt, Producer, Route, Vehicle
from apps.operations.services.bulk import HEADERS
from .factories import make_user


def upload(kind, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = {"producers": "Productores", "vehicles": "Vehículos", "routes": "Rutas", "receipts": "Recolecciones"}[kind]
    sheet.append(HEADERS[kind])
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile("carga.xlsx", out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class BulkTests(TestCase):
    def setUp(self):
        self.admin = make_user(username="bulkadmin", role=User.Role.ADMIN)
        self.supervisor = make_user(username="bulksupervisor", role=User.Role.SUPERVISOR)
        self.viewer = make_user(username="bulkviewer", role=User.Role.VIEWER)

    def post(self, kind, rows):
        return self.client.post(reverse("operations:bulk_import", kwargs={"kind": kind}), {"file": upload(kind, rows)})

    def test_routes_and_vehicles_are_upserted_and_exported(self):
        self.client.force_login(self.supervisor)
        self.post("routes", [["Ruta norte", "Primera zona"], ["Ruta sur", "Segunda zona"]])
        self.assertEqual(Route.objects.count(), 2)
        self.post("vehicles", [["abc123", "Ruta norte", "Ana", "1001", "300123", 1200, "Activo"]])
        vehicle = Vehicle.objects.get(plate="ABC123")
        self.assertEqual(vehicle.route.name, "RUTA NORTE")
        self.assertEqual(vehicle.driver_name, "Ana")
        self.post("vehicles", [["abc123", "Ruta sur", "Bea", "1002", "300456", 1500, "Cambio"]])
        self.assertEqual(Vehicle.objects.count(), 1)
        vehicle.refresh_from_db()
        self.assertEqual(vehicle.route.name, "RUTA SUR")
        response = self.client.get(reverse("operations:bulk_export", kwargs={"kind": "vehicles"}))
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        sheet = load_workbook(io.BytesIO(response.content), read_only=True)["Vehículos"]
        self.assertEqual(sheet["A1"].value, "Nombre / Placa del vehículo")
        self.assertEqual([sheet["A2"].value, sheet["B2"].value, sheet["C2"].value], ["ABC123", "RUTA SUR", "Bea"])

    def test_row_errors_prevent_partial_write(self):
        self.client.force_login(self.supervisor)
        response = self.post("producers", [["Cédula", "100", "María", "Finca", "", "Sogamoso", "Vereda 1", ""], ["Cédula", "", "Sin documento", "", "", "", "", ""]])
        self.assertContains(response, "Fila 3")
        self.assertEqual(Producer.objects.count(), 0)

    def test_permissions_on_both_import_modules(self):
        self.client.force_login(self.viewer)
        for kind in ("routes", "producers", "vehicles", "receipts"):
            self.assertEqual(self.client.get(reverse("operations:bulk_import", kwargs={"kind": kind})).status_code, 403)
        self.client.force_login(self.supervisor)
        self.assertEqual(self.client.get(reverse("operations:bulk_import", kwargs={"kind": "receipts"})).status_code, 403)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("operations:bulk_import", kwargs={"kind": "receipts"})).status_code, 200)

    def test_producer_import_updates_by_document(self):
        self.client.force_login(self.admin)
        row = ["NIT", "900467326-5", "VillaTech", "Finca A", "300", "Sogamoso", "Primavera", "Nota"]
        self.post("producers", [row])
        self.post("producers", [[*row[:2], "VillaTech nueva", *row[3:]]])
        self.assertEqual(Producer.objects.count(), 1)
        self.assertEqual(Producer.objects.get().name, "VillaTech nueva")

    def test_receipt_requires_real_time_and_existing_catalogs(self):
        self.client.force_login(self.admin)
        self.post("producers", [["Cédula", "100", "María", "Finca", "", "Sogamoso", "", ""]])
        self.post("vehicles", [["ABC123", "Ruta norte", "Ana", "", "", 1200, ""]])
        row = ["Ruta norte", "R-001", "100", "ABC123", "2026-09-24", "AM", 50, "Confirmado", "", "", "", ""]
        response = self.post("receipts", [row])
        self.assertContains(response, "HH:MM")
        self.assertEqual(MilkReceipt.objects.count(), 0)
        row[5] = "08:30"
        self.post("receipts", [row])
        receipt = MilkReceipt.objects.get()
        self.assertEqual(receipt.source, MilkReceipt.Source.IMPORT)
        self.assertEqual(receipt.source_reference, "R-001")
        response = self.post("receipts", [row])
        self.assertContains(response, "ya registrado")
        self.assertEqual(MilkReceipt.objects.count(), 1)

    def test_many_receipts_can_share_vehicle_and_producer(self):
        self.client.force_login(self.admin)
        self.post("producers", [["Otro", "PRUEBA-0001", "Productor prueba", "", "", "", "", "PRUEBA"]])
        self.post("vehicles", [["TST001", "PRUEBA NORTE", "", "", "", 1500, "PRUEBA"]])
        base = ["PRUEBA NORTE", "", "PRUEBA-0001", "TST001", "2026-09-20", "06:20", 85, "Confirmado", "", "", "", "PRUEBA"]
        rows = []
        for number in range(1, 4):
            row = base.copy()
            row[1] = f"PRUEBA-{number:03d}"
            rows.append(row)
        response = self.post("receipts", rows)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MilkReceipt.objects.count(), 3)
